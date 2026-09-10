from __future__ import annotations

import hashlib
import re
from collections.abc import Callable, Iterable
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .models import SourceChunk


SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp"}


class UnsupportedDocumentError(ValueError):
    pass


def discover_documents(input_path: Path) -> list[Path]:
    if input_path.is_file():
        candidates = [input_path]
    elif input_path.is_dir():
        candidates = sorted(path for path in input_path.rglob("*") if path.is_file())
    else:
        raise FileNotFoundError(f"Input does not exist: {input_path}")

    supported = [path for path in candidates if path.suffix.lower() in SUPPORTED_EXTENSIONS]
    if not supported:
        raise UnsupportedDocumentError(
            f"No supported documents found. Expected one of: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    return supported


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\x00", "")).strip()


def _markdown_sections(text: str) -> list[tuple[str, str]]:
    sections: list[tuple[str, list[str]]] = []
    heading = "Document"
    body: list[str] = []
    for line in text.splitlines():
        match = re.match(r"^#{1,6}\s+(.+?)\s*$", line)
        if match:
            if any(part.strip() for part in body):
                sections.append((heading, body))
            heading = _clean(match.group(1))
            body = []
        else:
            body.append(line)
    if any(part.strip() for part in body):
        sections.append((heading, body))
    if not sections and text.strip():
        sections.append((heading, [text]))
    return [(name, "\n".join(lines).strip()) for name, lines in sections]


def _text_sections(path: Path) -> list[tuple[str, str]]:
    return _markdown_sections(path.read_text(encoding="utf-8-sig"))


def _pdf_sections(path: Path) -> list[tuple[str, str]]:
    reader = PdfReader(str(path))
    return [
        (f"Page {number}", text)
        for number, page in enumerate(reader.pages, start=1)
        if (text := _clean(page.extract_text() or ""))
    ]


def _docx_sections(path: Path) -> list[tuple[str, str]]:
    document = Document(str(path))
    sections: list[tuple[str, list[str]]] = []
    heading = "Document"
    body: list[str] = []

    for paragraph in document.paragraphs:
        text = _clean(paragraph.text)
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            if body:
                sections.append((heading, body))
            heading, body = text, []
        else:
            body.append(text)

    table_rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            values = [_clean(cell.text) for cell in row.cells]
            table_rows.append(" | ".join(values))
    body.extend(table_rows)
    if body:
        sections.append((heading, body))
    return [(name, "\n\n".join(lines)) for name, lines in sections]


def _split_long_paragraph(text: str, max_chars: int) -> list[str]:
    sentences = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", text)
    parts: list[str] = []
    current = ""
    for sentence in sentences:
        if current and len(current) + len(sentence) + 1 > max_chars:
            parts.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        parts.append(current)
    return parts


def semantic_chunks(
    document: str,
    sections: Iterable[tuple[str, str]],
    max_chars: int = 1_600,
    overlap_paragraphs: int = 1,
) -> list[SourceChunk]:
    """Chunk on headings and paragraph/sentence boundaries, never raw token counts."""

    chunks: list[SourceChunk] = []
    ordinal = 0
    for section, text in sections:
        raw_paragraphs = [part.strip() for part in re.split(r"\n\s*\n|\n(?=[-*]\s)", text) if part.strip()]
        paragraphs: list[str] = []
        for paragraph in raw_paragraphs:
            paragraphs.extend(_split_long_paragraph(paragraph, max_chars))

        cursor = 0
        while cursor < len(paragraphs):
            selected: list[str] = []
            size = 0
            next_cursor = cursor
            while next_cursor < len(paragraphs):
                paragraph = paragraphs[next_cursor]
                addition = len(paragraph) + (2 if selected else 0)
                if selected and size + addition > max_chars:
                    break
                selected.append(paragraph)
                size += addition
                next_cursor += 1
            if not selected:
                selected = [paragraphs[cursor][:max_chars]]
                next_cursor = cursor + 1

            chunk_text = "\n\n".join(selected)
            digest = hashlib.sha1(
                f"{document}|{section}|{ordinal}|{chunk_text}".encode("utf-8")
            ).hexdigest()[:12]
            chunks.append(
                SourceChunk(
                    id=f"CHK-{digest.upper()}",
                    document=document,
                    section=section,
                    ordinal=ordinal,
                    text=chunk_text,
                )
            )
            ordinal += 1
            if next_cursor >= len(paragraphs):
                break
            cursor = max(cursor + 1, next_cursor - overlap_paragraphs)
    return chunks


class DocumentIngestor:
    def __init__(
        self,
        vision_describer: Callable[[Path], str],
        max_chunk_chars: int = 1_600,
        overlap_paragraphs: int = 1,
    ) -> None:
        self.vision_describer = vision_describer
        self.max_chunk_chars = max_chunk_chars
        self.overlap_paragraphs = overlap_paragraphs

    def ingest(self, input_path: Path) -> list[SourceChunk]:
        all_chunks: list[SourceChunk] = []
        for path in discover_documents(input_path):
            suffix = path.suffix.lower()
            if suffix in {".md", ".txt"}:
                sections = _text_sections(path)
            elif suffix == ".pdf":
                sections = _pdf_sections(path)
            elif suffix == ".docx":
                sections = _docx_sections(path)
            elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
                sections = [("Vision extraction", self.vision_describer(path))]
            else:  # guarded by discovery, retained as a defensive check
                raise UnsupportedDocumentError(f"Unsupported document: {path}")
            all_chunks.extend(
                semantic_chunks(
                    document=path.name,
                    sections=sections,
                    max_chars=self.max_chunk_chars,
                    overlap_paragraphs=self.overlap_paragraphs,
                )
            )
        if not all_chunks:
            raise ValueError("Documents were readable but contained no extractable text")
        return all_chunks
